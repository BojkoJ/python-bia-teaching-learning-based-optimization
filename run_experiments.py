"""
Spuštění experimentů pro všechny algoritmy (DE, PSO, SOMA, FA, TLBO)
na všech 9 testovacích funkcích s 30 opakováními a generování Excel souboru s výsledky.
"""
import numpy as np
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from main import (
    differential_evolution,
    particle_swarm_optimization,
    soma_all_to_one,
    firefly_algorithm,
    teaching_learning_based_optimization,
    sphere,
    schwefel,
    rosenbrock,
    rastrigin,
    griewank,
    levy,
    michalewicz,
    zakharov,
    ackley,
    get_default_bounds
)


# Konfigurace experimentů podle zadání
D = 30  # počet dimenzí
NP = 30  # velikost populace
MAX_OFE = 3000  # maximální počet vyhodnocení cílové funkce
NUM_EXPERIMENTS = 30  # počet experimentů

# Seznam testovacích funkcí
TEST_FUNCTIONS = [
    ("Sphere", sphere),
    ("Schwefel", schwefel),
    ("Rosenbrock", rosenbrock),
    ("Rastrigin", rastrigin),
    ("Griewank", griewank),
    ("Levy", levy),
    ("Michalewicz", michalewicz),
    ("Zakharov", zakharov),
    ("Ackley", ackley),
]


def run_single_experiment(algorithm_name, algorithm_func, objective, bounds, seed):
    """
    Spustí jeden experiment s daným algoritmem na dané funkci.
    
    Parametry:
        algorithm_name: název algoritmu (pro výpis)
        algorithm_func: funkce algoritmu
        objective: cílová funkce
        bounds: hranice pro dimenzi
        seed: náhodný seed
    
    Návrat:
        best_fitness: nejlepší nalezená hodnota fitness
    """
    try:
        if algorithm_name == "DE":
            # Differential Evolution
            _, best_fitness = algorithm_func(
                objective=objective,
                bounds=bounds,
                NP=NP,
                F=0.5,
                CR=0.5,
                max_gen=MAX_OFE // NP,  # MAX_OFE = max_gen * NP
                seed=seed,
                visualize=False
            )
        elif algorithm_name == "PSO":
            # Particle Swarm Optimization
            _, best_fitness = algorithm_func(
                objective=objective,
                bounds=bounds,
                num_particles=NP,
                w=0.7,
                c1=2.0,
                c2=2.0,
                max_iter=MAX_OFE // NP,  # MAX_OFE = max_iter * num_particles
                seed=seed,
                visualize=False
            )
        elif algorithm_name == "SOMA":
            # Self-Organizing Migrating Algorithm
            _, best_fitness = algorithm_func(
                objective=objective,
                bounds=bounds,
                pop_size=NP,
                path_length=3.0,
                step=0.11,
                prt=0.3,
                max_migrations=MAX_OFE // (NP * int(3.0 / 0.11)),  # Přibližný výpočet
                seed=seed,
                visualize=False
            )
        elif algorithm_name == "FA":
            # Firefly Algorithm
            _, best_fitness = algorithm_func(
                objective=objective,
                bounds=bounds,
                pop_size=NP,
                beta_0=1.0,
                alpha=0.3,
                max_generations=MAX_OFE // (NP * NP),  # FA má O(N^2) vyhodnocení na generaci
                seed=seed,
                visualize=False
            )
        elif algorithm_name == "TLBO":
            # Teaching-Learning Based Optimization
            _, best_fitness = algorithm_func(
                objective=objective,
                bounds=bounds,
                NP=NP,
                max_OFE=MAX_OFE,
                seed=seed,
                visualize=False
            )
        else:
            raise ValueError(f"Neznámý algoritmus: {algorithm_name}")
        
        return best_fitness
    except Exception as e:
        print(f"  CHYBA v {algorithm_name}: {e}")
        return float('inf')


def run_all_experiments():
    """
    Spustí všechny experimenty pro všechny algoritmy a funkce.
    
    Návrat:
        results: slovník s výsledky
                 results[func_name][algorithm_name] = list of 30 best fitness values
    """
    # Inicializujeme strukturu pro uložení výsledků
    results = {}
    
    # Algoritmy k testování
    algorithms = [
        ("DE", differential_evolution),
        ("PSO", particle_swarm_optimization),
        ("SOMA", soma_all_to_one),
        ("FA", firefly_algorithm),
        ("TLBO", teaching_learning_based_optimization),
    ]
    
    # Pro každou testovací funkci
    for func_name, func in TEST_FUNCTIONS:
        print("=" * 80)
        print(f"Spouštím experimenty pro funkci: {func_name}")
        print("=" * 80)
        
        # Získáme hranice pro D dimenzí
        bounds = get_default_bounds(func_name, D)
        
        # Inicializujeme slovník pro tuto funkci
        results[func_name] = {}
        
        # Pro každý algoritmus
        for algo_name, algo_func in algorithms:
            print(f"\n  Algoritmus: {algo_name}")
            
            # Inicializujeme seznam pro uložení výsledků
            algo_results = []
            
            # Spustíme NUM_EXPERIMENTS experimentů
            for exp_num in range(1, NUM_EXPERIMENTS + 1):
                # Seed pro reprodukovatelnost (různý pro každý experiment)
                seed = exp_num * 100
                
                # Spustíme experiment
                best_fitness = run_single_experiment(
                    algorithm_name=algo_name,
                    algorithm_func=algo_func,
                    objective=func,
                    bounds=bounds,
                    seed=seed
                )
                
                # Uložíme výsledek
                algo_results.append(best_fitness)
                
                # Výpis průběhu
                if exp_num % 10 == 0 or exp_num == 1:
                    print(f"    Experiment {exp_num:2d}/30: f = {best_fitness:.6e}")
            
            # Uložíme všechny výsledky pro tento algoritmus
            results[func_name][algo_name] = algo_results
            
            # Vypočítáme a vypíšeme statistiky
            mean = np.mean(algo_results)
            std = np.std(algo_results)
            print(f"    Mean: {mean:.6e}, Std: {std:.6e}")
    
    print("\n" + "=" * 80)
    print("VŠECHNY EXPERIMENTY DOKONČENY")
    print("=" * 80)
    
    return results


def create_excel_file(results, filename="results.xlsx"):
    """
    Vytvoří Excel soubor s 9 tabulkami (jedna pro každou funkci).
    
    Parametry:
        results: slovník s výsledky z run_all_experiments()
        filename: název výstupního souboru
    """
    print(f"\nVytvářím Excel soubor: {filename}")
    
    # Vytvoříme nový workbook
    wb = Workbook()
    
    # Odstraníme defaultní sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    
    # Algoritmy (sloupce v tabulce)
    algorithms = ["DE", "PSO", "SOMA", "FA", "TLBO"]
    
    # Pro každou funkci vytvoříme nový sheet
    for func_name in results.keys():
        print(f"  Vytvářím tabulku pro: {func_name}")
        
        # Vytvoříme nový sheet
        ws = wb.create_sheet(title=func_name)
        
        # Header (hlavička tabulky)
        ws.append([""] + algorithms)
        
        # Styling hlavičky
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for col_idx in range(1, len(algorithms) + 2):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Data - 30 experimentů
        for exp_num in range(1, NUM_EXPERIMENTS + 1):
            row_data = [f"Experiment {exp_num}"]
            
            for algo_name in algorithms:
                # Získáme výsledek tohoto experimentu
                value = results[func_name][algo_name][exp_num - 1]
                row_data.append(value)
            
            ws.append(row_data)
        
        # Prázdný řádek před statistikami
        ws.append([])
        
        # Mean a Std dev řádek
        mean_row = ["Mean"]
        std_row = ["Std dev"]
        
        for algo_name in algorithms:
            values = results[func_name][algo_name]
            mean_val = np.mean(values)
            std_val = np.std(values)
            mean_row.append(mean_val)
            std_row.append(std_val)
        
        ws.append(mean_row)
        ws.append(std_row)
        
        # Styling pro Mean/Std řádky
        mean_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        mean_font = Font(bold=True)
        for col_idx in range(1, len(algorithms) + 2):
            # Mean řádek
            cell = ws.cell(row=NUM_EXPERIMENTS + 3, column=col_idx)
            cell.fill = mean_fill
            cell.font = mean_font
            cell.alignment = Alignment(horizontal="center")
            
            # Std řádek
            cell = ws.cell(row=NUM_EXPERIMENTS + 4, column=col_idx)
            cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
            cell.font = mean_font
            cell.alignment = Alignment(horizontal="center")
        
        # Nastavíme šířky sloupců
        ws.column_dimensions['A'].width = 15
        for col_idx in range(2, len(algorithms) + 2):
            ws.column_dimensions[chr(64 + col_idx)].width = 25  # Širší sloupce pro čitatelnost
        
        # Formátování čísel (normální desetinná čísla místo vědecké notace)
        # Použijeme 12 desetinných míst aby se vešla i velmi malá čísla
        for row_idx in range(2, NUM_EXPERIMENTS + 5):
            for col_idx in range(2, len(algorithms) + 2):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None and isinstance(cell.value, (int, float)):
                    cell.number_format = '0.000000000000'  # 12 desetinných míst
    
    # Uložíme soubor
    wb.save(filename)
    print(f"\n✓ Excel soubor vytvořen: {filename}")


if __name__ == "__main__":
    print("=" * 80)
    print("SPUŠTĚNÍ EXPERIMENTŮ PRO VŠECHNY ALGORITMY")
    print("=" * 80)
    print(f"Konfigurace:")
    print(f"  Dimenze (D): {D}")
    print(f"  Velikost populace (NP): {NP}")
    print(f"  Max. vyhodnocení funkce (MAX_OFE): {MAX_OFE}")
    print(f"  Počet experimentů: {NUM_EXPERIMENTS}")
    print(f"  Testovací funkce: {len(TEST_FUNCTIONS)}")
    print()
    
    # Spustíme všechny experimenty
    results = run_all_experiments()
    
    # Vytvoříme Excel soubor
    create_excel_file(results, filename="tlbo_comparison_results.xlsx")
    
    print("\n" + "=" * 80)
    print("HOTOVO!")
    print("=" * 80)
